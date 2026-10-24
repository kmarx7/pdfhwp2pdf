import asyncio
from dataclasses import replace
import hashlib
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
import pytest

from app.config import Settings
from app.domain.errors import DomainError
from app.domain.models import ConversionOptions, DocumentAnalysis
from app.engines.pdf.pdf2docx_engine import Pdf2DocxEngine
from app.infrastructure.workspace import JobWorkspace
from app.main import create_app
from app.quality.validator import QualityValidator
from app.services.analyzer import DocumentAnalyzer
from app.services.job_service import JobService
from tests.conftest import submit, wait_job


def test_engine_page_exception_is_not_ignored(monkeypatch, tmp_path, samples):
    from pdf2docx import Converter
    options_seen = {}

    def fail(self, output, **kwargs):
        options_seen.update(kwargs)
        Path(output).write_bytes(b"partial file")
        raise RuntimeError("page failure")

    monkeypatch.setattr(Converter, "convert", fail)
    analysis = DocumentAnalyzer().analyze(samples["text"])
    with pytest.raises(DomainError) as error:
        Pdf2DocxEngine().convert(samples["text"], analysis, ConversionOptions(), tmp_path, lambda event: None)
    assert error.value.code == "ENGINE_FAILED"
    assert options_seen["ignore_page_error"] is False


@pytest.mark.parametrize("kind", ["image_only", "hidden_text", "missing_text", "missing_table"])
def test_validator_rejects_fake_or_incomplete_results(tmp_path, kind):
    from docx.oxml import OxmlElement
    document = Document()
    if kind != "image_only":
        run = document.add_paragraph().add_run("Some content" if kind != "missing_text" else "Some")
        if kind == "hidden_text":
            run._r.get_or_add_rPr().append(OxmlElement("w:vanish"))
    target = tmp_path / "bad.docx"
    document.save(target)
    analysis = DocumentAnalysis(1, [(595, 842)], ["Some content"], [], 1 if kind == "missing_table" else 0)
    with pytest.raises(DomainError) as error:
        QualityValidator().validate(target, analysis)
    assert error.value.code == "VALIDATION_FAILED"


@pytest.mark.parametrize(("setting", "value", "code"), [
    ("timeout_seconds", 0.001, "CONVERSION_TIMEOUT"),
    ("max_memory_bytes", 1, "RESOURCE_LIMIT"),
    ("max_pages", 1, "PAGE_LIMIT_EXCEEDED"),
])
def test_supervisor_limits_fail_and_clean_partial_output(tmp_path, samples, setting, value, code):
    settings = replace(Settings(), workspace=tmp_path / "limited", allowed_hosts=("testserver",), **{setting: value})
    with TestClient(create_app(settings)) as client:
        client.headers["X-CSRF-Token"] = client.get("/api/v1/session").json()["csrf_token"]
        response = submit(client, samples["multipage"])
        job = wait_job(client, response.json()["job_id"])
        assert job["state"] == "failed"
        assert job["error"]["code"] == code
        assert not (client.app.state.jobs.space(job["job_id"]).path / "output").exists()
        assert client.app.state.jobs.active is None


def test_symlink_workspace_refuses_access_and_deletion(tmp_path):
    from uuid import uuid4
    root = tmp_path / "jobs"
    root.mkdir()
    external = tmp_path / "original"
    external.mkdir()
    original = external / "important.pdf"
    original.write_bytes(b"original")
    key = str(uuid4())
    (root / key).symlink_to(external, target_is_directory=True)
    with pytest.raises(ValueError):
        JobWorkspace(root, key)
    assert original.read_bytes() == b"original"


def test_restart_recovers_interrupted_job_without_requeue(tmp_path, samples):
    from starlette.datastructures import UploadFile

    async def scenario():
        settings = replace(Settings(), workspace=tmp_path / "restart")
        first = JobService(settings)
        first.root.mkdir()
        job = await first.create(UploadFile(BytesIO(samples["text"].read_bytes()), filename="original.pdf"), "balanced", "off")
        key = job["job_id"]
        second = JobService(settings)
        await second.start()
        try:
            assert second.jobs[key]["state"] == "failed"
            assert second.jobs[key]["error"]["code"] == "INTERRUPTED"
            assert second.queue.empty()
            assert not (second.space(key).path / "input").exists()
        finally:
            await second.close()
    asyncio.run(scenario())


def test_tampered_completed_result_rejected_after_restart(tmp_path, samples):
    settings = replace(Settings(), workspace=tmp_path / "tamper", allowed_hosts=("testserver",))
    with TestClient(create_app(settings)) as client:
        client.headers["X-CSRF-Token"] = client.get("/api/v1/session").json()["csrf_token"]
        job = wait_job(client, submit(client, samples["text"]).json()["job_id"])
        assert job["state"] == "succeeded"
        path = client.app.state.jobs.space(job["job_id"]).path / "output/result.docx"
        path.write_bytes(b"tampered")
    with TestClient(create_app(settings)) as client:
        recovered = client.get(job["links"]["self"]).json()
        assert recovered["state"] == "failed"
        assert recovered["error"]["code"] == "VALIDATION_FAILED"


def test_second_server_cannot_share_workspace(tmp_path):
    settings = replace(Settings(), workspace=tmp_path / "single", allowed_hosts=("testserver",))
    with TestClient(create_app(settings)):
        with pytest.raises(RuntimeError, match="단일 서버"):
            with TestClient(create_app(settings)):
                pass
