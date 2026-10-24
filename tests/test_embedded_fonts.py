from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from lxml import etree
import pytest

from app.domain.errors import DomainError
from app.postprocess.embedded_fonts import embed_local_fonts, embedding_allowed, obfuscate, W, R, P
from scripts.convert_reference_pdf import convert


def make_font(path: Path, bold=False, flags=8, family="Fixture Sans", ps_name="FixtureSans"):
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder([".notdef", "A"])
    builder.setupCharacterMap({65: "A"})
    glyphs = {}
    for name in (".notdef", "A"):
        pen = TTGlyphPen(None)
        pen.moveTo((100, 0))
        pen.lineTo((500, 0))
        pen.lineTo((300, 700))
        pen.closePath()
        glyphs[name] = pen.glyph()
    builder.setupGlyf(glyphs)
    builder.setupHorizontalMetrics({name: (600, 0) for name in glyphs})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({"familyName": family, "styleName": "Bold" if bold else "Regular",
                            "uniqueFontIdentifier": "fixture-bold" if bold else "fixture-regular",
                            "fullName": family + (" Bold" if bold else " Regular"),
                            "psName": ps_name + ("-Bold" if bold else "")})
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200,
                     fsType=flags, fsSelection=32 if bold else 64)
    builder.setupPost()
    builder.setupMaxp()
    builder.save(path)


@pytest.fixture
def fonts(tmp_path):
    directory = tmp_path / "fonts"
    directory.mkdir()
    make_font(directory / "regular.ttf")
    make_font(directory / "bold.ttf", bold=True)
    return directory


def source_doc(path):
    document = Document()
    document.add_paragraph().add_run("AAA").font.name = "FixtureSans"
    document.save(path)


@pytest.mark.parametrize(("flags", "allowed"), [(0, True), (8, True), (256, True), (264, True),
    (2, False), (4, False), (512, False), (10, False), (32768, False)])
def test_editable_font_embedding_permission_flags(flags, allowed):
    assert embedding_allowed(flags) is allowed


def test_font_embedding_preserves_content_and_round_trips_font_bytes(tmp_path, fonts):
    source, target = tmp_path / "source.docx", tmp_path / "result.docx"
    source_doc(source)
    before = source.read_bytes()
    report = embed_local_fonts(source, target, [fonts])
    assert source.read_bytes() == before
    assert report["embedded_font_count"] == 2
    with ZipFile(source) as original, ZipFile(target) as result:
        assert result.read("word/document.xml") == original.read("word/document.xml")
        table = etree.fromstring(result.read("word/fontTable.xml"))
        relationships = etree.fromstring(result.read("word/_rels/fontTable.xml.rels"))
        by_id = {node.get("Id"): node.get("Target") for node in relationships}
        font = next(node for node in table if node.get(W + "name") == "FixtureSans")
        for tag, expected in (("embedRegular", "regular.ttf"), ("embedBold", "bold.ttf")):
            reference = font.find(W + tag)
            encoded = result.read("word/" + by_id[reference.get(R + "id")])
            decoded = obfuscate(encoded, reference.get(W + "fontKey"))
            assert decoded == (fonts / expected).read_bytes()
            assert encoded[:32] != decoded[:32] and encoded[32:] == decoded[32:]
    assert Document(target).paragraphs[0].text == "AAA"
    with pytest.raises(FileExistsError):
        embed_local_fonts(source, target, [fonts])


def test_missing_or_denied_font_never_creates_success_output(tmp_path, fonts):
    source, target = tmp_path / "source.docx", tmp_path / "result.docx"
    source_doc(source)
    with pytest.raises(DomainError, match="FONT_NOT_FOUND"):
        embed_local_fonts(source, target, [])
    assert not target.exists()
    make_font(fonts / "regular.ttf", flags=4)
    with pytest.raises(DomainError, match="FONT_EMBEDDING_DENIED"):
        embed_local_fonts(source, target, [fonts])
    assert not target.exists()


def test_reference_pipeline_refuses_overwrite_and_cleans_failed_work(tmp_path, samples):
    target = tmp_path / "result.docx"
    target.write_bytes(b"user result")
    with pytest.raises(FileExistsError):
        convert(samples["text"], target, [])
    assert target.read_bytes() == b"user result"
    other = tmp_path / "failed.docx"
    with pytest.raises(DomainError, match="FONT_NOT_FOUND"):
        convert(samples["text"], other, [])
    assert not other.exists() and not other.with_suffix(".quality.json").exists()
    assert not list(tmp_path.glob(".native-pdf-*"))
