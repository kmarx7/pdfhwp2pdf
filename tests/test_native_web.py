"""공개·합성 fixture만 사용한 native 웹 경로의 실제 엔진 회귀 검사."""
from dataclasses import replace
import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
import pymupdf
import pytest

from app.config import Settings
from app.domain.errors import DomainError
from app.main import create_app
from app.services.native_references import NativeReferences
from tests.conftest import wait_job
from tests.test_embedded_fonts import make_font

HWP = Path(__file__).parent / "fixtures/hwp/hello_world.hwp"


@pytest.fixture
def native_client(tmp_path):
    fonts = tmp_path / "fonts"
    fonts.mkdir()
    make_font(fonts / "regular.ttf", family="Helvetica", ps_name="Helvetica")
    make_font(fonts / "bold.ttf", bold=True, family="Helvetica", ps_name="Helvetica")
    settings = replace(Settings(), workspace=tmp_path / "jobs", native_font_dirs=(fonts,),
                       allowed_hosts=("testserver",), cleanup_interval=0.1)
    with TestClient(create_app(settings)) as client:
        client.headers["X-CSRF-Token"] = client.get("/api/v1/session").json()["csrf_token"]
        yield client


@pytest.fixture
def hello_pdf(tmp_path):
    path = tmp_path / "hello.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "Hello World!")
        document.save(path)
    return path


def native_submit(client, source, reference=None):
    files = {"file": (source.name, source.read_bytes())}
    if reference:
        files["reference"] = (reference.name, reference.read_bytes(), "application/pdf")
    return client.post("/api/v1/native/jobs", files=files)


@pytest.mark.parametrize("kind", ["text", "table", "multipage"])
def test_native_pdf_real_engine_text_tables_fonts_download_and_cleanup(native_client, samples, kind):
    source = samples[kind]
    before = source.read_bytes()
    response = native_submit(native_client, source)
    assert response.status_code == 202, response.text
    job = wait_job(native_client, response.json()["job_id"])
    assert job["state"] == "succeeded", job
    assert job["engine"]["id"] == "native-pdf-embedded"
    download = native_client.get(job["links"]["download"])
    assert download.headers["content-disposition"] == 'attachment; filename="converted.docx"'
    assert download.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    document = Document(BytesIO(download.content))
    assert "Editable PDF conversion" in "".join(p.text for p in document.paragraphs)
    if kind == "table":
        assert "Alpha" in document.tables[0].cell(1, 0).text
    with ZipFile(BytesIO(download.content)) as package:
        assert b"embedRegular" in package.read("word/fontTable.xml")
        assert not any(name.startswith("word/media/") for name in package.namelist())
    quality = native_client.get(job["links"]["quality"]).json()
    assert quality["route"] == "native_reference_pdf_to_docx_with_fonts"
    assert quality["embedded_fonts"]["embedded_font_count"] == 2
    assert quality["metrics"]["character_coverage"]["value"] == 1
    assert source.read_bytes() == before
    space = native_client.app.state.jobs.space(job["job_id"]).path
    assert not (space / "input").exists() and not (space / "work").exists()


@pytest.mark.parametrize("cached", [False, True])
def test_native_hwp_uses_attached_or_hash_bound_reference(native_client, hello_pdf, cached):
    before = HWP.read_bytes()
    if cached:
        digest = native_client.app.state.jobs.references.register(HWP, hello_pdf)
        status = native_client.get(f"/api/v1/native/references/{digest}").json()
        assert status["available"] is True and status["pages"] == 1
        assert "path" not in status
    response = native_submit(native_client, HWP, None if cached else hello_pdf)
    assert response.status_code == 202, response.text
    job = wait_job(native_client, response.json()["job_id"])
    assert job["state"] == "succeeded", job
    assert job["format"] == "hwp" and job["mode"] == "native_embedded"
    assert job["engine"]["id"] == "native-pdf-embedded"
    quality = native_client.get(job["links"]["quality"]).json()
    assert quality["source_format"] == "hwp"
    assert quality["metrics"]["hwp_reference_text_coverage"]["value"] == 1
    assert quality["metrics"]["reference_pdf_page_count"]["value"] == 1
    assert quality["metrics"]["source_page_count"]["status"] == "not_measured"
    assert quality["input_sha256"] == hashlib.sha256(before).hexdigest()
    assert quality["reference_pdf_sha256"] == hashlib.sha256(hello_pdf.read_bytes()).hexdigest()
    document = Document(BytesIO(native_client.get(job["links"]["download"]).content))
    assert "Hello World!" in "".join(p.text for p in document.paragraphs)
    assert HWP.read_bytes() == before


def test_hwp_without_reference_never_silently_falls_back(native_client):
    response = native_submit(native_client, HWP)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "NATIVE_REFERENCE_REQUIRED"
    assert not native_client.app.state.jobs.jobs
    assert not list(native_client.app.state.jobs.root.glob("*/input"))


@pytest.mark.parametrize("kind,code", [("text", "REFERENCE_MISMATCH"), ("encrypted", "PASSWORD_PROTECTED"),
                                      ("invalid", "INVALID_DOCUMENT")])
def test_wrong_or_invalid_reference_has_no_download(native_client, samples, kind, code):
    response = native_submit(native_client, HWP, samples[kind])
    assert response.status_code == 202, response.text
    job = wait_job(native_client, response.json()["job_id"])
    assert job["state"] == "failed" and job["error"]["code"] == code, job
    assert job["links"]["download"] is None
    assert not (native_client.app.state.jobs.space(job["job_id"]).path / "output").exists()


def test_missing_font_fails_without_output(native_client, hello_pdf):
    for path in native_client.app.state.jobs.settings.native_font_dirs[0].iterdir():
        path.unlink()
    job = wait_job(native_client, native_submit(native_client, hello_pdf).json()["job_id"])
    assert job["state"] == "failed" and job["error"]["code"] == "FONT_NOT_FOUND", job
    assert job["links"]["download"] is None


def test_reference_cache_binding_expiry_and_integrity(tmp_path, hello_pdf):
    now = [100.0]
    references = NativeReferences(tmp_path / "references", ttl=10, clock=lambda: now[0])
    digest = references.register(HWP, hello_pdf)
    assert references.lookup(digest)[1]["expires_at"] == 110
    assert references.lookup("0" * 64) is None
    assert not list(references.root.glob(".register-*"))
    assert not list(references.root.rglob("source-model.json"))
    references.lookup(digest)[0].write_bytes(b"tampered")
    with pytest.raises(DomainError, match="REFERENCE_INVALID"):
        references.lookup(digest)
    now[0] = 111
    assert references.lookup(digest) is None and not references.entry(digest).exists()
    with pytest.raises(DomainError, match="INVALID_REQUEST"):
        references.lookup("../outside")
    references.entry(digest).symlink_to(tmp_path, target_is_directory=True)
    with pytest.raises(DomainError, match="REFERENCE_INVALID"):
        references.lookup(digest)


def test_partial_reference_is_rejected_and_registration_cleans_model(native_client, tmp_path):
    partial = tmp_path / "partial.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "Hello")
        document.save(partial)
    references = native_client.app.state.jobs.references
    with pytest.raises(DomainError, match="REFERENCE_MISMATCH"):
        references.register(HWP, partial)
    assert not list(references.root.iterdir())
    job = wait_job(native_client, native_submit(native_client, HWP, partial).json()["job_id"])
    assert job["state"] == "failed" and job["error"]["code"] == "REFERENCE_MISMATCH", job
    assert job["links"]["download"] is None


def test_native_request_validation_and_csrf(native_client, hello_pdf):
    assert native_client.post("/api/v1/native/jobs", files={"file": ("source.pdf", hello_pdf.read_bytes())},
                              headers={"Origin": "https://other.example"}).status_code == 403
    token = native_client.headers.pop("X-CSRF-Token")
    assert native_submit(native_client, hello_pdf).status_code == 403
    native_client.headers["X-CSRF-Token"] = token
    assert native_submit(native_client, hello_pdf, hello_pdf).json()["error"]["code"] == "INVALID_OPTIONS"
    assert native_client.post("/api/v1/native/jobs", files={"file": ("source.pdf", hello_pdf.read_bytes()),
                                                           "extra": ("extra.pdf", hello_pdf.read_bytes())}).status_code == 422
