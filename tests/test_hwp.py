import copy
import hashlib
import json
import struct
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from app.domain.errors import DomainError
from app.domain.models import ConversionOptions
from app.engines.hwp.cli import HwpCli
from app.engines.hwp.hwp_cli_engine import HwpCliEngine
from app.postprocess.docx import DocxPostProcessor
from app.postprocess.hwp import HwpDocxPostProcessor, column_widths
from app.postprocess.hwp_fills import resolve_fill
from app.quality.validator import QualityValidator, normalized
from app.services.hwp_analyzer import HwpAnalyzer, guard_hwp_container
from app.services.hwp_page_numbers import parse_page_number
from scripts.make_test_hwps import make_test_hwps
from tests.conftest import submit, wait_job


@pytest.fixture(scope="session")
def hwp_samples(tmp_path_factory):
    assert HwpCli().available(), "HWP 통합 테스트 전에 scripts/install_hwp_engine.py를 실행하세요."
    return make_test_hwps(tmp_path_factory.mktemp("synthetic-hwp"))


@pytest.mark.parametrize("kind", ["text", "table", "merged", "image", "sections", "colored"])
def test_hwp_real_engine_download_editable_objects(client, hwp_samples, kind):
    source = hwp_samples[kind]
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    response = submit(client, source)
    assert response.status_code == 202, response.text
    assert response.json()["format"] == "hwp"
    job = wait_job(client, response.json()["job_id"])
    assert job["state"] == "succeeded", job
    assert job["engine"] == {"id": "hwp-cli", "version": "1.3.1"}
    assert job["progress"]["total_pages"] is None
    download = client.get(job["links"]["download"])
    assert download.status_code == 200
    assert download.headers["content-disposition"] == f'attachment; filename="{source.stem}.docx"'
    assert download.headers["content-type"].endswith("wordprocessingml.document")
    document = Document(BytesIO(download.content))
    text = "".join(node.text or "" for node in document.element.findall(".//" + qn("w:t")))
    assert "한글HWP변환12345" in normalized(text)
    run = next(run for para in document.paragraphs for run in para.runs if "12345" in run.text)
    assert run.bold is True and run.font.size.pt == 18 and run.font.name == "Arial"
    assert str(run.font.color.rgb) == "112233"
    assert abs(document.sections[0].page_width.pt - 595.28) < 1
    if kind in {"table", "colored"}:
        assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [["품목", "금액"], ["사과", "100"], ["배", "200"]]
        assert abs(document.tables[0].columns[0].width.mm - 40) < 0.1
        assert abs(document.tables[0].columns[1].width.mm - 110) < 0.1
        if kind == "colored":
            assert document.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "FFFFCC"
    if kind == "merged":
        table = document.tables[0]
        assert table.cell(0, 0)._tc is table.cell(1, 0)._tc
        assert table.cell(0, 1)._tc is table.cell(0, 2)._tc
    if kind == "image":
        assert len(document.inline_shapes) == 1
        with ZipFile(BytesIO(download.content)) as package:
            assert len([name for name in package.namelist() if name.startswith("word/media/")]) == 1
    else:
        with ZipFile(BytesIO(download.content)) as package:
            assert not any(name.startswith("word/media/") for name in package.namelist())
    if kind == "sections":
        assert len(document.sections) == 2
        assert "두번째섹션의본문67890" in normalized(text)
    document.paragraphs[0].add_run("수정")
    if document.tables:
        document.tables[0].cell(0, 0).text = "셀 수정"
    saved = BytesIO()
    document.save(saved)
    saved.seek(0)
    assert "수정" in Document(saved).paragraphs[0].text
    report = client.get(job["links"]["quality"]).json()
    assert report["metrics"]["character_coverage"]["value"] == 1
    assert report["metrics"]["source_page_count"]["status"] == "not_measured"
    assert report["source_format"] == "hwp"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    workspace = client.app.state.jobs.space(job["job_id"]).path
    assert not (workspace / "input").exists() and not (workspace / "work").exists()


def test_independent_apache_hwp_fixture(client):
    source = Path(__file__).parent / "fixtures/hwp/hello_world.hwp"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == "283e16ee51834d75e22f413203549027a5d3fda39346302d9429e9732bb3a77c"
    job = wait_job(client, submit(client, source).json()["job_id"])
    assert job["state"] == "succeeded", job
    document = Document(BytesIO(client.get(job["links"]["download"]).content))
    assert "Hello" in "".join(paragraph.text for paragraph in document.paragraphs)


@pytest.mark.parametrize(("kind", "code"), [
    ("password_flag", "PASSWORD_PROTECTED"), ("distribution", "RESTRICTED_DOCUMENT"),
    ("drm", "RESTRICTED_DOCUMENT"), ("header", "UNSUPPORTED_FEATURE"),
])
def test_hwp_protection_and_unsupported_features_fail_without_output(client, hwp_samples, kind, code):
    job = wait_job(client, submit(client, hwp_samples[kind]).json()["job_id"])
    assert job["state"] == "failed", job
    assert job["error"]["code"] == code
    if kind == "header":
        assert job["error"]["details"]["feature_labels"] == ["머리말"]
        assert "감지된 미지원 요소: 머리말." in job["error"]["suggestion"]
    assert job["links"]["download"] is None
    assert client.get(job["links"]["self"] + "/download").status_code == 409
    assert not (client.app.state.jobs.space(job["job_id"]).path / "output").exists()


@pytest.mark.parametrize(("control", "label"), [("pghd", "쪽 감추기"), ("zzzz", "미지원 문서 요소")])
def test_hwp_unsupported_error_names_detected_feature(control, label):
    model = {"sections": [{"paragraphs": [{
        "chars": [{"Text": "본문"}, {"ExtCtrl": {"ctrl_index": 0}}],
        "controls": [{"Generic": {"ctrl_id": list(control.encode("ascii"))}}],
    }]}]}
    with pytest.raises(DomainError) as error:
        HwpAnalyzer().inspect_model(model)
    assert error.value.code == "UNSUPPORTED_FEATURE"
    assert error.value.details == {"features": [control], "feature_labels": [label]}
    assert f"감지된 미지원 요소: {label}." in error.value.suggestion
    assert "머리말·꼬리말" not in error.value.suggestion


def test_hwp_invalid_and_mismatched_files(client, hwp_samples, samples):
    assert submit(client, hwp_samples["text"], "wrong.pdf").json()["error"]["code"] == "FORMAT_MISMATCH"
    assert submit(client, samples["text"], "wrong.hwp").json()["error"]["code"] == "FORMAT_MISMATCH"
    assert client.post("/api/v1/jobs", files={"file": ("invalid.hwp", b"invalid")}).status_code == 400
    response = client.post("/api/v1/jobs", files={"file": ("broken.hwp", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"X" * 512)})
    job = wait_job(client, response.json()["job_id"])
    assert job["error"]["code"] == "INVALID_DOCUMENT"


def test_old_hwp_variant_is_explicitly_rejected(client):
    response = client.post("/api/v1/jobs", files={"file": ("old.hwp", b"HWP Document File V3.00" + b"X" * 512)})
    job = wait_job(client, response.json()["job_id"])
    assert job["error"]["code"] == "UNSUPPORTED_VARIANT"


def test_hwp_compression_guard_and_missing_engine(hwp_samples, monkeypatch, tmp_path):
    import app.services.hwp_analyzer as analyzer
    monkeypatch.setattr(analyzer, "MAX_STREAM", 300)
    with pytest.raises(DomainError) as error:
        guard_hwp_container(hwp_samples["text"])
    assert error.value.code == "ARCHIVE_LIMIT_EXCEEDED"
    assert not HwpCli(binary=tmp_path / "missing").available()


@pytest.mark.parametrize("kind", ["table", "merged", "image"])
def test_hwp_result_structure_loss_is_rejected(hwp_samples, tmp_path, kind):
    analysis = HwpAnalyzer().analyze(hwp_samples[kind], tmp_path)
    result = HwpCliEngine().convert(hwp_samples[kind], analysis, ConversionOptions(), tmp_path, lambda event: None)
    HwpDocxPostProcessor().process(result.docx_path, analysis)
    DocxPostProcessor().process(result.docx_path)
    QualityValidator().validate(result.docx_path, analysis)
    broken = copy.deepcopy(analysis)
    if kind == "image":
        broken.engine_data["image_hashes"] = ["wrong"]
    else:
        broken.engine_data["tables"][0]["cells"][0]["row_span"] += 1
    with pytest.raises(DomainError) as error:
        QualityValidator().validate(result.docx_path, broken)
    assert error.value.code == "VALIDATION_FAILED"


def test_hwp_engine_failure_does_not_publish_result(hwp_samples, tmp_path, monkeypatch):
    analysis = HwpAnalyzer().analyze(hwp_samples["text"], tmp_path)

    def fail(*args, **kwargs):
        raise DomainError("ENGINE_FAILED", "테스트 엔진 실패")

    monkeypatch.setattr(HwpCli, "run", fail)
    with pytest.raises(DomainError):
        HwpCliEngine().convert(hwp_samples["text"], analysis, ConversionOptions(), tmp_path, lambda event: None)
    assert not (tmp_path / "converted.tmp.docx").exists()


def test_hwp_binary_integrity_and_installer_rejects_invalid_archive(tmp_path):
    from scripts.install_hwp_engine import install
    fake = tmp_path / "hwp"
    fake.write_bytes(b"different binary")
    fake.chmod(0o755)
    assert not HwpCli(binary=fake).available()
    archive = tmp_path / "invalid.tar.gz"
    archive.write_bytes(b"invalid archive")
    with pytest.raises(ValueError, match="SHA-256"):
        install(archive)


@pytest.mark.parametrize("kind", ["page_numbers", "roman_numbers", "section_numbers", "inherited_numbers"])
def test_hwp_page_numbers_are_live_word_fields(client, hwp_samples, kind):
    source = hwp_samples[kind]
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    job = wait_job(client, submit(client, source).json()["job_id"])
    assert job["state"] == "succeeded", job
    document = Document(BytesIO(client.get(job["links"]["download"]).content))
    section = document.sections[0]
    part = section.header if kind == "roman_numbers" else section.footer
    fields = part._element.findall(".//" + qn("w:fldSimple"))
    assert len(fields) == 1
    field = fields[0]
    assert field.get(qn("w:dirty")) == "true"
    assert field.get(qn("w:instr")).strip() == ("PAGE \\* ROMAN" if kind == "roman_numbers" else "PAGE \\* ARABIC")
    paragraph = part.paragraphs[0]
    assert paragraph.alignment == (2 if kind == "roman_numbers" else 1)  # Word right/center
    assert paragraph.text == ("()" if kind == "roman_numbers" else "-  -")
    # 번호는 본문 숫자/이미지가 아닌 자동 필드이며 번호를 계산할 뷰어가 필요하다.
    assert field.find(qn("w:r")) is None
    number_type = section._sectPr.find(qn("w:pgNumType"))
    assert number_type.get(qn("w:start")) == ("4" if kind == "roman_numbers" else "1")
    if kind == "section_numbers":
        second = document.sections[1]
        assert not second.header.is_linked_to_previous and not second.footer.is_linked_to_previous
        assert not second.footer._element.findall(".//" + qn("w:fldSimple"))
        assert not second.header._element.findall(".//" + qn("w:fldSimple"))
    if kind == "inherited_numbers":
        second = document.sections[1]
        assert second.footer._element.find(".//" + qn("w:fldSimple")).get(qn("w:instr")).strip() == "PAGE \\* ARABIC"
        assert second._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) is None
    assert any("PAGE 필드" in warning for warning in job["warnings"])
    assert before == hashlib.sha256(source.read_bytes()).hexdigest()


@pytest.mark.parametrize(("position", "number_format", "data"), [
    (7, 0, None), (8, 0, None), (5, 1, None), (5, 99, None), (5, 0, "00"), (5, 0, "not-hex"),
])
def test_hwp_unsupported_page_number_forms_are_not_silently_replaced(position, number_format, data):
    payload = struct.pack("<I4H", (position << 8) | number_format, 0, 0, 0, 0).hex() if data is None else data
    with pytest.raises(DomainError) as error:
        parse_page_number({"data": payload})
    assert error.value.code == "UNSUPPORTED_FEATURE"
    assert error.value.details["features"] == ["pgnp"]
    assert error.value.details["reason"]


def test_hwp_mid_section_page_number_is_rejected(hwp_samples, tmp_path):
    model_path = tmp_path / "source.json"
    HwpCli().run(["convert", str(hwp_samples["page_numbers"]), "-o", str(model_path), "--embed-bin"], tmp_path)
    model = json.loads(model_path.read_text())
    first = model["sections"][0]["paragraphs"][0]
    control = next(control for control in first["controls"] if "Generic" in control and control["Generic"]["ctrl_id"] == list(b"pgnp"))
    # 첫 설정은 유지하되 본문 뒤에서 다른 위치를 지정한다.
    last = model["sections"][0]["paragraphs"][-1]
    changed = copy.deepcopy(control)
    changed["Generic"]["data"] = struct.pack("<I4H", 2 << 8, 0, 0, 0, 0).hex()
    index = len(last["controls"])
    last["controls"].append(changed)
    last["chars"].insert(0, {"ExtCtrl": {"ctrl_index": index}})
    with pytest.raises(DomainError) as error:
        HwpAnalyzer().inspect_model(model)
    assert "구역 본문 중간" in error.value.message


def test_hwp_missing_page_number_field_is_rejected(hwp_samples, tmp_path):
    source = hwp_samples["page_numbers"]
    analysis = HwpAnalyzer().analyze(source, tmp_path)
    result = HwpCliEngine().convert(source, analysis, ConversionOptions(), tmp_path, lambda event: None)
    HwpDocxPostProcessor().process(result.docx_path, analysis)
    QualityValidator().validate(result.docx_path, analysis)
    document = Document(result.docx_path)
    field = document.sections[0].footer._element.find(".//" + qn("w:fldSimple"))
    field.getparent().remove(field)
    document.save(result.docx_path)
    with pytest.raises(DomainError) as error:
        QualityValidator().validate(result.docx_path, analysis)
    assert error.value.code == "VALIDATION_FAILED"


def test_hwp_gradient_approximation_and_dotted_double_borders(client, hwp_samples):
    job = wait_job(client, submit(client, hwp_samples["decorated"]).json()["job_id"])
    assert job["state"] == "succeeded", job
    assert any("그라데이션 효과는 보존되지 않았습니다" in warning for warning in job["warnings"])
    doc = Document(BytesIO(client.get(job["links"]["download"]).content))
    properties = doc.tables[0].cell(0, 0)._tc.tcPr
    assert properties.find(qn("w:shd")).get(qn("w:fill")) == "EFF2FB"
    assert properties.find(qn("w:tcBorders")).find(qn("w:left")).get(qn("w:val")) == "dotted"
    assert properties.find(qn("w:tcBorders")).find(qn("w:top")).get(qn("w:val")) == "double"
    report = client.get(job["links"]["quality"]).json()
    assert report["warnings"] == job["warnings"]
    assert report["metrics"]["character_coverage"]["value"] == 1


def test_hwp_background_loss_is_rejected(hwp_samples, tmp_path):
    source = hwp_samples["decorated"]
    analysis = HwpAnalyzer().analyze(source, tmp_path)
    result = HwpCliEngine().convert(source, analysis, ConversionOptions(), tmp_path, lambda event: None)
    HwpDocxPostProcessor().process(result.docx_path, analysis)
    doc = Document(result.docx_path)
    properties = doc.tables[0].cell(0, 0)._tc.tcPr
    properties.remove(properties.find(qn("w:shd")))
    doc.save(result.docx_path)
    with pytest.raises(DomainError) as error:
        QualityValidator().validate(result.docx_path, analysis)
    assert error.value.code == "VALIDATION_FAILED"


@pytest.mark.parametrize("difference", [283, 301])
def test_hwp_small_row_width_difference_is_bounded_and_recorded(difference):
    table = {"cols": 2, "cells": [
        {"col": 0, "col_span": 2, "width": 30000 - difference},
        {"col": 0, "col_span": 1, "width": 10000},
        {"col": 1, "col_span": 1, "width": 20000},
    ]}
    if difference > 300:
        with pytest.raises(DomainError):
            column_widths(table)
    else:
        assert column_widths(table) == [10000, 20000]
        assert table["width_adjustment_hu"] == 283


def test_hwp_image_fill_and_unsupported_gradients_remain_rejected():
    for fill in ({"fill_type": 2}, {"fill_type": 4, "gradient": {"radial": True}},
                 {"fill_type": 4, "gradient": {"radial": False, "stops": [[0, 0], [0.5, 1], [1, 2]]}}):
        with pytest.raises(DomainError) as error:
            resolve_fill(fill)
        assert error.value.code == "UNSUPPORTED_FEATURE"
    assert resolve_fill({"fill_type": 1, "bg_color": 0xFFFFFFFF})["background"] is None


def test_hwp_empty_paragraph_line_spacing_and_explicit_page_break(client, hwp_samples):
    job = wait_job(client, submit(client, hwp_samples["layout"]).json()["job_id"])
    assert job["state"] == "succeeded", job
    doc = Document(BytesIO(client.get(job["links"]["download"]).content))
    assert len(doc.paragraphs) == 4  # 구역 설정 빈 문단 + 제목 + 원본 빈 문단 + 본문
    assert doc.paragraphs[2].text == ""
    assert doc.paragraphs[-1].paragraph_format.page_break_before is True
    assert abs(doc.paragraphs[-1].paragraph_format.line_spacing - 1.6) < 0.01
    assert "Editable text." in doc.paragraphs[-1].text
    # HWP 본문 영역은 상단 여백과 머리말 영역을 함께 포함한다.
    assert doc.sections[0].top_margin.mm > 18
    report = client.get(job["links"]["quality"]).json()
    assert report["metrics"]["hwp_layout_hints"]["value"]["explicit_page_breaks"] == 1
    assert report["metrics"]["hwp_layout_hints"]["visual_fidelity_verified"] is False


def test_hwp_row_height_is_minimum_and_cells_keep_vertical_alignment(client, hwp_samples):
    job = wait_job(client, submit(client, hwp_samples["table"]).json()["job_id"])
    assert job["state"] == "succeeded", job
    doc = Document(BytesIO(client.get(job["links"]["download"]).content))
    assert all(row.height is not None and row.height_rule == 1 for row in doc.tables[0].rows)
    assert doc.tables[0].cell(0, 0).vertical_alignment == 1
