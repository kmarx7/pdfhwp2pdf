import hashlib
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from app.main import DOCX_MIME
from app.quality.validator import normalized
from tests.conftest import submit, wait_job


@pytest.mark.parametrize("kind", ["text", "table", "korean", "multipage", "image"])
def test_real_engine_creates_editable_docx(client, samples, kind):
    source = samples[kind]
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    response = submit(client, source)
    assert response.status_code == 202, response.text
    assert response.headers["location"] == response.json()["links"]["self"]
    job = wait_job(client, response.json()["job_id"])
    assert job["state"] == "succeeded", job
    download = client.get(job["links"]["download"])
    assert download.status_code == 200
    assert download.headers["content-type"] == DOCX_MIME
    assert download.headers["content-disposition"] == 'attachment; filename="converted.docx"'
    assert download.headers["cache-control"] == "no-store"
    assert download.headers["x-content-type-options"] == "nosniff"
    document = Document(BytesIO(download.content))
    # WordprocessingML 본문이 실제 텍스트 노드로 존재한다.
    nodes = document.element.findall(".//" + qn("w:t"))
    assert "EditablePDFconversion" in normalized("".join(node.text or "" for node in nodes))
    with ZipFile(BytesIO(download.content)) as package:
        assert b"<w:t" in package.read("word/document.xml")
        assert package.testzip() is None
        if kind == "text":
            assert not any(name.startswith("word/media/") for name in package.namelist())
    if kind == "table":
        assert len(document.tables) == 1
        assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [["Item", "Amount"], ["Alpha", "100"], ["Beta", "200"]]
        assert len(document.tables[0].rows) == 3
    if kind == "korean":
        assert "한글문서변환테스트" in normalized("".join(node.text or "" for node in nodes))
    if kind == "image":
        assert len(document.inline_shapes) >= 1
    if kind == "multipage":
        assert "Secondlandscapepage67890." in normalized("".join(node.text or "" for node in nodes))
        assert any(section.page_width > section.page_height for section in document.sections)
    # 구조적 편집/저장·재열기. 실제 Word 앱 수동 편집 검증과는 구분한다.
    document.paragraphs[0].add_run(" Edited")
    if document.tables:
        document.tables[0].cell(1, 0).text = "Changed"
        document.tables[0].add_row()
    edited = BytesIO()
    document.save(edited)
    edited.seek(0)
    assert "Edited" in Document(edited).paragraphs[0].text
    quality = client.get(job["links"]["quality"]).json()
    assert quality["validation_status"] == "structural_pass"
    assert quality["metrics"]["character_coverage"]["value"] >= 0.99
    assert quality["manual_review"]["word_edit"] == "not_performed"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    space = client.app.state.jobs.space(job["job_id"]).path
    assert not (space / "input").exists()
    assert not (space / "work").exists()


@pytest.mark.parametrize(("kind", "code"), [
    ("invalid", "INVALID_DOCUMENT"), ("encrypted", "PASSWORD_PROTECTED"),
    ("owner_protected", "PASSWORD_PROTECTED"), ("scanned", "OCR_REQUIRED"),
    ("mixed", "OCR_REQUIRED"), ("ocr_layer", "OCR_REQUIRED"), ("blank", "NO_EDITABLE_TEXT"),
])
def test_bad_or_unsupported_pdf_fails_without_download(client, samples, kind, code):
    response = submit(client, samples[kind])
    assert response.status_code == 202
    job = wait_job(client, response.json()["job_id"])
    assert job["state"] == "failed", job
    assert job["error"]["code"] == code, job
    assert job["links"]["download"] is None
    assert client.get(job["links"]["self"] + "/download").status_code == 409
    space = client.app.state.jobs.space(job["job_id"]).path
    assert not (space / "input").exists()
    assert not (space / "output").exists()
    if kind == "mixed":
        assert job["error"]["details"]["pages"] == [2]


def test_untrusted_filename_never_used_as_path_or_header(client, samples):
    response = submit(client, samples["text"], '../../한글"; injected.pdf')
    assert response.status_code == 202
    job = wait_job(client, response.json()["job_id"])
    download = client.get(job["links"]["download"])
    assert download.headers["content-disposition"] == 'attachment; filename="converted.docx"'
    meta = (client.app.state.jobs.space(job["job_id"]).path / "meta.json").read_text()
    assert "injected" not in meta and "한글" not in meta
