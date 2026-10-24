"""기존 DOCX 텍스트·셀을 유지하면서 HWP 표의 열 치수와 기본 테두리/배경을 복원한다."""
from collections import deque
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.domain.errors import DomainError
from app.domain.models import DocumentAnalysis


def column_widths(table: dict) -> list[int]:
    count = table["cols"]
    graph = [[] for _ in range(count + 1)]
    for cell in table["cells"]:
        start, end, width = cell["col"], cell["col"] + cell["col_span"], cell["width"]
        if not 0 <= start < end <= count or width <= 0:
            raise DomainError("UNSUPPORTED_FEATURE", "HWP 표 열 치수를 확인할 수 없습니다.")
        graph[start].append((end, width))
        graph[end].append((start, -width))
    boundaries = {0: 0}
    queue = deque([0])
    while queue:
        start = queue.popleft()
        for end, distance in graph[start]:
            position = boundaries[start] + distance
            if end not in boundaries:
                boundaries[end] = position
                queue.append(end)
            elif abs(boundaries[end] - position) > 10:
                raise DomainError("UNSUPPORTED_FEATURE", "HWP 표의 열 치수가 일관되지 않아 변환을 중단했습니다.")
    if count not in boundaries:
        raise DomainError("UNSUPPORTED_FEATURE", "HWP 표 전체 너비를 확인할 수 없습니다.")
    # 모든 행에서 병합된 경계는 실제 선이 없다. 알려진 양 끝 치수 사이만 분할한다.
    known = sorted(boundaries)
    for left, right in zip(known, known[1:]):
        for index in range(left + 1, right):
            boundaries[index] = round(boundaries[left] + (boundaries[right] - boundaries[left]) * (index - left) / (right - left))
    widths = [boundaries[index + 1] - boundaries[index] for index in range(count)]
    if min(widths) <= 0:
        raise DomainError("UNSUPPORTED_FEATURE", "HWP 표 열 너비가 올바르지 않습니다.")
    return widths


def colorref(value: int) -> str:
    return f"{value & 255:02X}{(value >> 8) & 255:02X}{(value >> 16) & 255:02X}"


def replace_property(parent, name: str, attributes: dict, before: str | None = None):
    existing = parent.find(qn(name))
    position = None
    if existing is not None:
        position = list(parent).index(existing)
        parent.remove(existing)
    element = OxmlElement(name)
    for key, value in attributes.items():
        element.set(qn(key), str(value))
    sibling = parent.find(qn(before)) if before else None
    if position is not None:
        parent.insert(position, element)
    elif sibling is None:
        parent.append(element)
    else:
        parent.insert(list(parent).index(sibling), element)
    return element


class HwpDocxPostProcessor:
    def process(self, path: Path, analysis: DocumentAnalysis) -> None:
        document = Document(path)
        tables = document.element.findall(".//" + qn("w:tbl"))
        sources = analysis.engine_data["tables"]
        if len(tables) != len(sources):
            raise DomainError("VALIDATION_FAILED", "변환 결과의 HWP 표 개수가 다릅니다.")
        fills = analysis.engine_data["border_fills"]
        for target, source in zip(tables, sources):
            widths = column_widths(source)
            source["column_widths"] = widths
            grid = target.find(qn("w:tblGrid"))
            if grid is None or len(grid) != len(widths):
                raise DomainError("VALIDATION_FAILED", "변환 결과의 HWP 표 열 구조가 다릅니다.")
            for element, width in zip(grid, widths):
                element.set(qn("w:w"), str(round(width / 5)))
            properties = target.find(qn("w:tblPr"))
            replace_property(properties, "w:tblW", {"w:w": round(sum(widths) / 5), "w:type": "dxa"}, "w:tblBorders")
            for border in properties.findall(qn("w:tblBorders")):
                properties.remove(border)
            cells_by_position = {}
            for original in source["cells"]:
                for row in range(original["row"], original["row"] + original["row_span"]):
                    cells_by_position[(row, original["col"])] = original
            for row_index, row in enumerate(target.findall(qn("w:tr"))):
                column = 0
                for cell in row.findall(qn("w:tc")):
                    cell_properties = cell.find(qn("w:tcPr"))
                    span = cell_properties.find(qn("w:gridSpan"))
                    colspan = int(span.get(qn("w:val"))) if span is not None else 1
                    width = round(sum(widths[column:column + colspan]) / 5)
                    replace_property(cell_properties, "w:tcW", {"w:w": width, "w:type": "dxa"}, "w:gridSpan")
                    original = cells_by_position.get((row_index, column))
                    if original:
                        self.cell_format(cell_properties, original, fills, row_index)
                    column += colspan
        self.page_numbers(document, analysis)
        document.save(path)

    @staticmethod
    def page_numbers(document, analysis: DocumentAnalysis):
        placements = analysis.engine_data.get("page_numbers", [])
        if not any(placements):
            return
        if len(placements) != len(document.sections):
            raise DomainError("VALIDATION_FAILED", "쪽 번호의 구역 구조가 다릅니다.")
        document.settings.odd_and_even_pages_header_footer = False
        for index, (section, placement) in enumerate(zip(document.sections, placements)):
            # 이전 구역 번호가 번호 없는 구역으로 연결되지 않도록 각 구역을 분리한다.
            section.header.is_linked_to_previous = False
            section.footer.is_linked_to_previous = False
            section.different_first_page_header_footer = False
            attributes = {"w:fmt": placement["format"] if placement else "decimal"}
            if index == 0:
                attributes["w:start"] = analysis.engine_data["page_start"]
            replace_property(section._sectPr, "w:pgNumType", attributes, "w:cols")
            if not placement or not placement["position"]:
                continue
            position = placement["position"]
            paragraph = (section.header if position <= 3 else section.footer).paragraphs[0]
            paragraph.alignment = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT][(position - 1) % 3]
            paragraph.add_run(placement["prefix"])
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), f' PAGE \\* {placement["field_format"]} ')
            field.set(qn("w:dirty"), "true")
            # 캐시 숫자를 고정 본문으로 만들지 않는다. 실제 번호는 Word가 계산한다.
            paragraph._p.append(field)
            paragraph.add_run(placement["suffix"])
        settings = document.settings.element
        replace_property(settings, "w:updateFields", {"w:val": "true"}, "w:compat")

    @staticmethod
    def cell_format(properties, cell: dict, fills: list[dict], row: int):
        index = cell["border_fill"] - 1
        if not 0 <= index < len(fills):
            raise DomainError("UNSUPPORTED_FEATURE", "HWP 표의 테두리/배경 참조가 올바르지 않습니다.")
        fill = fills[index]
        if fill.get("fill_type", 0) not in (0, 1):
            raise DomainError("UNSUPPORTED_FEATURE", "그라데이션·그림 채우기 표 배경은 아직 지원하지 않습니다.")
        thickness = [0.1, 0.12, 0.15, 0.2, 0.25, 0.3, 0.4, 0.5, 0.6, 0.7, 1, 1.5, 2, 3, 4, 5]
        borders = OxmlElement("w:tcBorders")
        sides = dict(zip(("left", "right", "top", "bottom"), fill["sides"]))
        for name in ("top", "left", "bottom", "right"):
            side = sides[name]
            if side["line_type"] not in (0, 1):
                raise DomainError("UNSUPPORTED_FEATURE", "현재 단일 실선/선 없음 표 테두리만 보존합니다.")
            continuing = (name == "top" and row > cell["row"]) or (name == "bottom" and row < cell["row"] + cell["row_span"] - 1)
            border = OxmlElement("w:" + name)
            border.set(qn("w:val"), "nil" if not side["line_type"] or continuing else "single")
            border.set(qn("w:color"), colorref(side["color"]))
            border.set(qn("w:sz"), str(max(2, round(thickness[min(side["width"], 15)] * 72 / 25.4 * 8))))
            borders.append(border)
        for existing in properties.findall(qn("w:tcBorders")):
            properties.remove(existing)
        properties.append(borders)
        if fill.get("bg_color") is not None:
            replace_property(properties, "w:shd", {"w:val": "clear", "w:color": "auto", "w:fill": colorref(fill["bg_color"])})
        margins = replace_property(properties, "w:tcMar", {})
        values = dict(zip(("left", "right", "top", "bottom"), cell["margins"]))
        for name in ("top", "left", "bottom", "right"):
            value = values[name]
            replace_property(margins, "w:" + name, {"w:w": max(0, round(value / 5)), "w:type": "dxa"})
